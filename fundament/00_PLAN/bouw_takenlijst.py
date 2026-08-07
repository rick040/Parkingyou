#!/usr/bin/env python3
"""Bouwt Takenlijst_en_Mijlpalen.xlsx uit Projectplan_v1.0.md.

Het projectplan is de enige bron van waarheid voor de taken. Door de xlsx
hieruit te genereren kunnen de twee niet uit elkaar lopen.
"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

PLAN = '/home/user/Parkingyou/fundament/00_PLAN/Projectplan_v1.0.md'
UIT = '/home/user/Parkingyou/fundament/00_PLAN/Takenlijst_en_Mijlpalen.xlsx'

FASENAAM = {
    '0': '0 · Fundament',
    '1': '1 · Abonnementen',
    '2': '2 · Eerste koppeling',
    '3': '3 · Uitrollen',
    '4': '4 · Klantenservice',
    '5': '5 · Automatiseren',
}
FASEPERIODE = {
    '0': 'week 1–4', '1': 'week 4–8', '2': 'week 6–14',
    '3': 'maand 4–7', '4': 'maand 6–9', '5': 'maand 9–12',
}
# Welke taak sluit welke mijlpaal af (uit het projectplan, §5 en §6)
MIJLPAAL_NA = {
    '0.9': 'M0', '0.44': 'M1', '0.52': 'M2',
    '1.13': 'M4', '1.18': 'M3',
    '2.14': 'M5', '2.23': 'M6',
    '3.8': 'M7', '3.21': 'M8',
    '4.15': 'M9',
    '5.12': 'M10', '5.24': 'M11',
}
MIJLPALEN = [
    ('M0', '0', 'Besluiten genomen, kritiek pad gestart', 'Directie', '0.9',
     'Vier leveranciersaanvragen verstuurd, met datum in de Connector Registry'),
    ('M1', '0', 'Registers live', 'D2 Operatie', '0.44',
     '37 locaties compleet in NocoDB; v_datakwaliteit leeg'),
    ('M2', '0', 'Documentlaag live, oude drive bevroren', 'Directie', '0.52',
     'Nieuwe SharePoint in gebruik; schrijven in de oude drive wordt geweigerd'),
    ('M3', '1', 'Abonnementen gemigreerd', 'D3 Klantenservice', '1.18',
     'Aantallen kloppen met Excel; steekproef van tien correct; Excel bevroren'),
    ('M4', '1', 'Kentekenwijziging werkt end-to-end', 'D3 Klantenservice', '1.13',
     'Formulier werkt; mutatie automatisch vastgelegd; werklijst loopt leeg'),
    ('M5', '2', 'Pilot-parkeersysteem levert data', 'D1 Financiën', '2.14',
     'Dagelijkse run draait vanzelf; tweede import verdubbelt niet; fout geeft alarm'),
    ('M6', '2', 'Dashboard per locatie', 'D1 + D2', '2.23',
     'Dwarsdoorsnede klopt; advieslocaties tellen niet mee; prognose naar rato'),
    ('M7', '3', 'Alle platformen en AFAS gekoppeld', 'D1 Financiën', '3.8',
     'Nul exploitatielocaties zonder data over de laatste 7 dagen'),
    ('M8', '3', 'Eigenaarsrapportages automatisch', 'Directie', '3.21',
     'Gegenereerd, juiste naam en map, verzonden aan de primaire contactpartij'),
    ('M9', '4', 'Kanaal gekoppeld en kennisbank gevuld', 'D3 Klantenservice', '4.15',
     'Klantvragen landen in de datalaag; top 20 vragen heeft een actueel artikel'),
    ('M10', '5', 'AI-klantenservice live', 'D3 + Directie', '5.12',
     'Minimaal 15 van 20 testvragen correct; nul foute antwoorden zonder escalatie'),
    ('M11', '5', 'Signaallus actief', 'Directie', '5.24',
     'Signaal → oorzaakcheck → campagne → ROI, met een uitgelokte afwijking getest'),
]

# ---------------------------------------------------------------- taken lezen
src = open(PLAN, encoding='utf-8').read()

# Kopjes van de blokken binnen een fase (#### 0A. ...)
blokken = {}
huidig_blok = ''
taken = []
for regel in src.split('\n'):
    kop = re.match(r'^####\s+(\d[A-Z])\.\s+(.+?)\s*$', regel)
    if kop:
        huidig_blok = f"{kop.group(1)} {kop.group(2)}"
        continue
    if re.match(r'^###\s+Fase', regel):
        huidig_blok = ''
        continue
    m = re.match(r'^\|\s*(\d+\.\d+)\s*\|([^|]+)\|([^|]+)\|([^|]+)\|([^|]*)\|\s*$', regel)
    if m:
        tid, taak, eig, uren, afh = (x.strip() for x in m.groups())
        try:
            u = float(uren.replace(',', '.'))
        except ValueError:
            u = 0.0
        taken.append({
            'id': tid, 'fase': tid.split('.')[0], 'blok': huidig_blok,
            'taak': re.sub(r'\*\*(.+?)\*\*', r'\1', taak),
            'eigenaar': eig, 'uren': u,
            'afh': '' if afh == '—' else afh,
            'mijlpaal': MIJLPAAL_NA.get(tid, ''),
        })

assert len(taken) == 175, f'verwacht 175 taken, gevonden {len(taken)}'

# ---------------------------------------------------------------- opmaak
KOP = PatternFill('solid', fgColor='1F3864')
FASE_KLEUR = {
    '0': 'DDEBF7', '1': 'E2EFDA', '2': 'FFF2CC',
    '3': 'FCE4D6', '4': 'E4DFEC', '5': 'D9E1F2',
}
MIJLPAAL_FILL = PatternFill('solid', fgColor='FFD966')
WIT = Font(color='FFFFFF', bold=True, size=11)
RAND = Border(*[Side(style='thin', color='BFBFBF')] * 4)

wb = Workbook()

# ============================================================ blad: Takenlijst
ws = wb.active
ws.title = 'Takenlijst'
kolommen = [
    ('Nr', 8), ('Fase', 20), ('Blok', 30), ('Taak', 74), ('Eigenaar', 11),
    ('Uren', 8), ('Hangt af van', 14), ('Mijlpaal', 10),
    ('Status', 15), ('Gereed op', 12), ('Opmerking', 34),
]
for i, (naam, breedte) in enumerate(kolommen, 1):
    c = ws.cell(row=1, column=i, value=naam)
    c.fill, c.font = KOP, WIT
    c.alignment = Alignment(vertical='center')
    ws.column_dimensions[get_column_letter(i)].width = breedte
ws.row_dimensions[1].height = 22

for r, t in enumerate(taken, start=2):
    waarden = [
        t['id'], FASENAAM[t['fase']], t['blok'], t['taak'], t['eigenaar'],
        t['uren'], t['afh'], t['mijlpaal'], 'Niet gestart', None, None,
    ]
    for i, v in enumerate(waarden, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = RAND
        c.fill = PatternFill('solid', fgColor=FASE_KLEUR[t['fase']])
        c.alignment = Alignment(vertical='top', wrap_text=(i == 4))
        if i == 6:
            c.number_format = '0.00'
        if i == 10:
            c.number_format = 'dd-mm-yyyy'
    if t['mijlpaal']:
        for i in (1, 8):
            ws.cell(row=r, column=i).fill = MIJLPAAL_FILL
            ws.cell(row=r, column=i).font = Font(bold=True)

laatste = len(taken) + 1
ws.auto_filter.ref = f'A1:K{laatste}'
ws.freeze_panes = 'A2'

status = DataValidation(
    type='list',
    formula1='"Niet gestart,Bezig,Wacht op derden,Geblokkeerd,Gereed,Vervallen"',
    allow_blank=True)
status.add(f'I2:I{laatste}')
ws.add_data_validation(status)

ws.conditional_formatting.add(
    f'I2:I{laatste}',
    CellIsRule(operator='equal', formula=['"Gereed"'],
               fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(color='006100')))
ws.conditional_formatting.add(
    f'I2:I{laatste}',
    CellIsRule(operator='equal', formula=['"Geblokkeerd"'],
               fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(color='9C0006')))
ws.conditional_formatting.add(
    f'I2:I{laatste}',
    CellIsRule(operator='equal', formula=['"Wacht op derden"'],
               fill=PatternFill('solid', fgColor='FFEB9C'), font=Font(color='9C5700')))

# ============================================================ blad: Mijlpalen
wm = wb.create_sheet('Mijlpalen')
mk = [('Mijlpaal', 10), ('Fase', 20), ('Wat er dan staat', 44),
      ('Tekent af', 20), ('Sluit af na taak', 16), ('Testcriterium', 66),
      ('Voortgang t/m', 13), ('Status', 15), ('Afgetekend op', 14)]
for i, (naam, breedte) in enumerate(mk, 1):
    c = wm.cell(row=1, column=i, value=naam)
    c.fill, c.font = KOP, WIT
    wm.column_dimensions[get_column_letter(i)].width = breedte
wm.row_dimensions[1].height = 22

rij_van_taak = {t['id']: i + 2 for i, t in enumerate(taken)}

for r, (code, fase, wat, tekent, na, test) in enumerate(MIJLPALEN, start=2):
    # Voortgang tot en met de taak die deze mijlpaal afsluit — niet over de hele
    # fase, anders tonen alle mijlpalen binnen één fase hetzelfde percentage.
    fase_taken = [t for t in taken if t['fase'] == fase]
    eerste_rij = rij_van_taak[fase_taken[0]['id']]
    laatste_rij = rij_van_taak[na]
    aantal = laatste_rij - eerste_rij + 1
    voortgang = (f'=COUNTIF(Takenlijst!$I${eerste_rij}:$I${laatste_rij},"Gereed")'
                 f'/{aantal}')
    for i, v in enumerate([code, FASENAAM[fase], wat, tekent, na, test,
                           voortgang, 'Niet gehaald', None], 1):
        c = wm.cell(row=r, column=i, value=v)
        c.border = RAND
        c.fill = PatternFill('solid', fgColor=FASE_KLEUR[fase])
        c.alignment = Alignment(vertical='top', wrap_text=(i in (3, 6)))
        if i == 1:
            c.font = Font(bold=True)
        if i == 7:
            c.number_format = '0%'
        if i == 9:
            c.number_format = 'dd-mm-yyyy'

mv = DataValidation(type='list', formula1='"Niet gehaald,In test,Gehaald"', allow_blank=True)
mv.add(f'H2:H{len(MIJLPALEN)+1}')
wm.add_data_validation(mv)
wm.conditional_formatting.add(
    f'H2:H{len(MIJLPALEN)+1}',
    CellIsRule(operator='equal', formula=['"Gehaald"'],
               fill=PatternFill('solid', fgColor='C6EFCE'), font=Font(color='006100')))
wm.freeze_panes = 'A2'

# ============================================================ blad: Overzicht
wo = wb.create_sheet('Overzicht', 0)
wo['A1'] = 'ParkingYou Fundament — voortgang'
wo['A1'].font = Font(size=16, bold=True, color='1F3864')
wo['A2'] = 'Gegenereerd uit Projectplan_v1.0.md. Werk de kolom Status bij op het blad Takenlijst.'
wo['A2'].font = Font(italic=True, color='595959')

kop = ['Fase', 'Periode', 'Taken', 'Gereed', 'Voortgang',
       'Uren R+C', 'Uren domein', 'Uren directie', 'Uren totaal']
for i, naam in enumerate(kop, 1):
    c = wo.cell(row=4, column=i, value=naam)
    c.fill, c.font = KOP, WIT
for i, b in enumerate([22, 13, 8, 9, 11, 11, 13, 13, 12], 1):
    wo.column_dimensions[get_column_letter(i)].width = b
wo.column_dimensions['A'].width = 22

def uren(fase, groep):
    tot = 0.0
    for t in taken:
        if t['fase'] != fase:
            continue
        e = t['eigenaar']
        is_rc = e in ('R', 'C', 'R+C', 'C+R')
        is_dir = 'DIR' in e and '+' not in e
        if groep == 'rc' and is_rc:
            tot += t['uren']
        elif groep == 'dir' and is_dir:
            tot += t['uren']
        elif groep == 'dom' and not is_rc and not is_dir:
            tot += t['uren']
    return tot

rij = 5
for fase in sorted(FASENAAM):
    ft = [t for t in taken if t['fase'] == fase]
    eerste = taken.index(ft[0]) + 2
    laatste_t = taken.index(ft[-1]) + 2
    wo.cell(row=rij, column=1, value=FASENAAM[fase])
    wo.cell(row=rij, column=2, value=FASEPERIODE[fase])
    wo.cell(row=rij, column=3, value=len(ft))
    wo.cell(row=rij, column=4,
            value=f'=COUNTIF(Takenlijst!I{eerste}:I{laatste_t},"Gereed")')
    c = wo.cell(row=rij, column=5, value=f'=IF(C{rij}=0,0,D{rij}/C{rij})')
    c.number_format = '0%'
    wo.cell(row=rij, column=6, value=uren(fase, 'rc')).number_format = '0.00'
    wo.cell(row=rij, column=7, value=uren(fase, 'dom')).number_format = '0.00'
    wo.cell(row=rij, column=8, value=uren(fase, 'dir')).number_format = '0.00'
    wo.cell(row=rij, column=9, value=f'=SUM(F{rij}:H{rij})').number_format = '0.00'
    for i in range(1, 10):
        wo.cell(row=rij, column=i).fill = PatternFill('solid', fgColor=FASE_KLEUR[fase])
        wo.cell(row=rij, column=i).border = RAND
    rij += 1

wo.cell(row=rij, column=1, value='Totaal').font = Font(bold=True)
for col in 'CDFGHI':
    wo[f'{col}{rij}'] = f'=SUM({col}5:{col}{rij-1})'
    wo[f'{col}{rij}'].font = Font(bold=True)
    if col not in 'CD':
        wo[f'{col}{rij}'].number_format = '0.00'
wo[f'E{rij}'] = f'=IF(C{rij}=0,0,D{rij}/C{rij})'
wo[f'E{rij}'].number_format = '0%'
wo[f'E{rij}'].font = Font(bold=True)
for i in range(1, 10):
    wo.cell(row=rij, column=i).border = RAND

r2 = rij + 2
wo.cell(row=r2, column=1, value='Mijlpalen').font = Font(size=13, bold=True, color='1F3864')
wo.cell(row=r2 + 1, column=1, value='Gehaald')
wo.cell(row=r2 + 1, column=3,
        value=f'=COUNTIF(Mijlpalen!H2:H{len(MIJLPALEN)+1},"Gehaald") & " van {len(MIJLPALEN)}"')

r3 = r2 + 3
wo.cell(row=r3, column=1, value='Let op').font = Font(size=13, bold=True, color='C00000')
for i, tekst in enumerate([
    'De leveranciersaanvragen (taken 0.3 t/m 0.6) bepalen de doorlooptijd van het hele project.',
    'Zet ze in week 1 op "Wacht op derden" en bewaak ze — niets in fase 2 en 3 kan zonder.',
    'Een mijlpaal is pas gehaald als het testscript in Mijlpaal_Testrapporten.md klopt,',
    'niet als de taken zijn afgevinkt. De aftekenaar is nooit de bouwer.',
]):
    wo.cell(row=r3 + 1 + i, column=1, value=tekst)

wb.save(UIT)
print(f'{UIT} geschreven — {len(taken)} taken, {len(MIJLPALEN)} mijlpalen')
